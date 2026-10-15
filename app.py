import re
from io import BytesIO
from datetime import datetime

import fitz
import pandas as pd
import streamlit as st
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

st.set_page_config(
    page_title="HireLens AI",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
)

SKILL_ALIASES = {
    "python": ["python"], "java": ["java"], "c++": ["c++"], "c#": ["c#"],
    "sql": ["sql", "mysql", "postgresql", "sqlite"],
    "machine learning": ["machine learning", "scikit-learn", "sklearn"],
    "deep learning": ["deep learning", "pytorch", "tensorflow", "keras"],
    "data analysis": ["data analysis", "data analytics"], "pandas": ["pandas"],
    "numpy": ["numpy"], "excel": ["excel", "microsoft excel"],
    "power bi": ["power bi"], "tableau": ["tableau"], "html": ["html"],
    "css": ["css"], "javascript": ["javascript"], "typescript": ["typescript"],
    "react": ["react", "react.js"], "node.js": ["node.js", "nodejs"],
    "django": ["django"], "flask": ["flask"], "streamlit": ["streamlit"],
    "git": ["git", "github", "gitlab"], "aws": ["aws", "amazon web services"],
    "azure": ["azure", "microsoft azure"], "docker": ["docker"],
    "kubernetes": ["kubernetes"], "rest api": ["rest api", "restful api"],
    "communication": ["communication", "verbal communication"],
    "project management": ["project management"],
    "leadership": ["leadership", "team leadership"],
    "problem solving": ["problem solving", "analytical thinking"],
    "linux": ["linux"], "networking": ["networking", "computer networks"],
    "cybersecurity": ["cybersecurity", "information security"],
    "mongodb": ["mongodb"], "firebase": ["firebase"], "fastapi": ["fastapi"],
    "c": [" c ", "c programming"], "data structures": ["data structures", "dsa"],
    "algorithms": ["algorithms"], "oop": ["object oriented programming", "oop"],
    "testing": ["software testing", "unit testing", "pytest"],
    "agile": ["agile", "scrum"], "figma": ["figma"], "ui/ux": ["ui/ux", "user experience"],
}


def norm(text):
    return re.sub(r"\s+", " ", str(text).lower()).strip()


def extract_text(uploaded_file):
    data = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()
    if filename.endswith(".pdf"):
        with fitz.open(stream=data, filetype="pdf") as pdf:
            return "\n".join(page.get_text() for page in pdf)
    if filename.endswith(".docx"):
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError("Unsupported file type")


def detect_skills(text):
    text = norm(text)
    found = set()
    for skill, aliases in SKILL_ALIASES.items():
        for alias in aliases:
            # Word-boundary-like matching avoids matching 'java' inside 'javascript'.
            pattern = r"(?<!\w)" + re.escape(alias.strip().lower()) + r"(?!\w)"
            if re.search(pattern, text):
                found.add(skill)
                break
    return found


def parse_skill_list(value):
    items = [x.strip().lower() for x in value.split(",") if x.strip()]
    # Map known aliases back to canonical names.
    canonical = []
    for item in items:
        matched = None
        for skill, aliases in SKILL_ALIASES.items():
            if item == skill or item in [a.lower() for a in aliases]:
                matched = skill
                break
        canonical.append(matched or item)
    return list(dict.fromkeys(canonical))


def analyze(job_description, resumes, must_have, preferred, text_weight):
    names = list(resumes.keys())
    texts = [resumes[name] for name in names]
    vectorizer = TfidfVectorizer(
        stop_words="english", ngram_range=(1, 2),
        sublinear_tf=True, max_features=25000,
    )
    matrix = vectorizer.fit_transform([job_description] + texts)
    similarities = cosine_similarity(matrix[0:1], matrix[1:]).flatten()

    required = set(must_have)
    preferred_set = set(preferred) - required
    rows = []
    for idx, name in enumerate(names):
        found = detect_skills(texts[idx])
        req_match = required & found
        req_missing = required - found
        pref_match = preferred_set & found
        req_coverage = len(req_match) / len(required) if required else None
        pref_coverage = len(pref_match) / len(preferred_set) if preferred_set else None

        # Score blends text relevance, must-have coverage, and preferred skills.
        # If a criterion group is empty, its weight is redistributed.
        components = [(text_weight, float(similarities[idx]))]
        if required:
            components.append(((1 - text_weight) * (0.75 if preferred_set else 1.0), req_coverage))
        if preferred_set:
            components.append(((1 - text_weight) * 0.25, pref_coverage))
        total_weight = sum(w for w, _ in components)
        score = 100 * sum(w * val for w, val in components) / total_weight if total_weight else 0

        # Explainable triage guidance only; not an automated hiring decision.
        if req_missing:
            recommendation = "Review gaps before proceeding"
            recommendation_reason = "One or more must-have skills were not detected; verify the resume because parsing or wording can miss skills."
        elif score >= 70:
            recommendation = "Recommend for interview review"
            recommendation_reason = "Strong document overlap and all listed must-have skills were detected. Verify experience and evidence in an interview."
        else:
            recommendation = "Needs human review"
            recommendation_reason = "The document match is below the suggested review threshold; inspect relevant experience and resume evidence before deciding."

        rows.append({
            "Candidate": name,
            "AI recommendation": recommendation,
            "Recommendation rationale": recommendation_reason,
            "Match score": round(float(score), 1),
            "Text relevance": round(float(similarities[idx]) * 100, 1),
            "Must-have coverage": round(req_coverage * 100, 1) if req_coverage is not None else None,
            "Preferred coverage": round(pref_coverage * 100, 1) if pref_coverage is not None else None,
            "Matching must-haves": ", ".join(sorted(req_match)) or "None detected",
            "Must-haves to verify": ", ".join(sorted(req_missing)) or "None detected",
            "Preferred skills found": ", ".join(sorted(pref_match)) or "None detected",
            "_text": texts[idx],
            "_found": found,
            "_missing": req_missing,
        })
    return pd.DataFrame(rows).sort_values(
        ["Match score", "Text relevance"], ascending=False
    ).reset_index(drop=True)


def demo_resumes():
    return {
        "Aarav_Shah_DEMO.txt": (
            "Software engineering graduate. Projects using Python, SQL, pandas, "
            "machine learning and Flask. Built a student performance prediction "
            "project and REST API. Familiar with Git, data analysis, DSA and communication."
        ),
        "Meera_Rao_DEMO.txt": (
            "Computer science student with Python, scikit-learn, pandas, NumPy, "
            "deep learning and TensorFlow. Built a classification model and deployed "
            "a Streamlit dashboard. Used GitHub, SQL and testing."
        ),
        "Kabir_Nair_DEMO.txt": (
            "Frontend developer with HTML, CSS, JavaScript, React and TypeScript. "
            "Built responsive web applications. Uses Git, REST APIs, UI/UX and problem solving."
        ),
        "Diya_Menon_DEMO.txt": (
            "Data analyst intern experienced with Excel, SQL, Power BI, Tableau, "
            "Python and pandas. Created dashboards, cleaned datasets and presented insights."
        ),
        "Ishaan_Verma_DEMO.txt": (
            "Backend developer using Python, FastAPI, Django, PostgreSQL, Docker, "
            "AWS and REST APIs. Built services with testing, Git and CI workflows."
        ),
    }


def apply_style():
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700;800&family=Space+Grotesk:wght@500;600;700&display=swap');

    :root {
        --ink: #18213b;
        --muted: #68738c;
        --purple: #7257f5;
        --cyan: #35c9d6;
        --pink: #ff6eaa;
        --line: rgba(112, 126, 165, .22);
    }
    html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
    .stApp {
        background:
          radial-gradient(ellipse at 5% 0%, rgba(91, 113, 255, .12), transparent 32%),
          radial-gradient(ellipse at 100% 12%, rgba(255, 110, 170, .10), transparent 28%),
          linear-gradient(180deg, #f8f9ff 0%, #f4f7fc 48%, #fbfaff 100%);
    }
    [data-testid="stHeader"] { background: rgba(248,249,255,.78); }
    .block-container { max-width: 1460px; padding-top: 1.3rem; padding-bottom: 1rem; }
    h1,h2,h3,h4 { font-family: 'Space Grotesk', sans-serif; letter-spacing: -.035em; color: #202945; }
    h2, h3 { margin-top: .5rem; }
    p, li, label { color: #34405d; }
    .hero {
        position: relative; overflow: hidden; isolation: isolate;
        padding: clamp(1.5rem, 3vw, 2.35rem);
        border-radius: 26px;
        background: linear-gradient(118deg, #17224a 0%, #3e3a96 48%, #7655e8 77%, #3fc4d0 125%);
        color: #fff; margin: .25rem 0 1.45rem 0;
        box-shadow: 0 18px 45px rgba(67, 69, 151, .20);
        border: 1px solid rgba(255,255,255,.18);
    }
    .hero:before {
        content: ""; position: absolute; z-index: -1; width: 260px; height: 260px;
        right: -65px; top: -105px; border-radius: 50%;
        background: radial-gradient(circle, rgba(255,255,255,.22), rgba(255,255,255,0) 70%);
    }
    .hero:after {
        content: ""; position: absolute; z-index: -1; width: 190px; height: 190px;
        right: 16%; bottom: -145px; border: 1px solid rgba(255,255,255,.22);
        border-radius: 50%; box-shadow: 0 0 0 22px rgba(255,255,255,.035), 0 0 0 44px rgba(255,255,255,.025);
    }
    .hero h1 { color: white !important; text-align: center; margin: .2rem 0 .75rem 0; font-size: clamp(2.15rem, 4.3vw, 3.1rem); line-height: 1.06; font-weight: 800; }
    .hero p { color: #e8eaff !important; margin: 0 auto; max-width: 800px; font-size: 1.02rem; line-height: 1.65; }
    .eyebrow {
        display: inline-flex; align-items: center; gap: .45rem;
        text-transform: uppercase; letter-spacing: .15em; font-size: .70rem;
        font-weight: 800; color: #d6faff; margin-bottom: .5rem;
        border: 1px solid rgba(255,255,255,.25); background: rgba(255,255,255,.10);
        border-radius: 999px; padding: .35rem .7rem;
    }
    .brandline { color: #dce3ff; font-size: .78rem; font-weight: 700; letter-spacing: .12em; margin-top: 1.05rem; }
    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, #171f43 0%, #25255e 56%, #303c79 100%);
        border-right: 1px solid rgba(255,255,255,.12);
    }
    [data-testid="stSidebar"] * { color: #f5f6ff; }
    [data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p { color: #c8d0f4; }
    [data-testid="stSidebar"] hr { border-color: rgba(255,255,255,.18); }
    [data-testid="stSidebar"] [data-testid="stRadio"] label { color: #f5f6ff; }
    [data-testid="stSidebar"] [data-testid="stToggle"] label { color: #f5f6ff; }
    [data-testid="stSidebar"] button { border: 1px solid rgba(255,255,255,.20); }
    div[data-testid="stMetric"] {
        background: rgba(255,255,255,.84); border: 1px solid var(--line);
        padding: .9rem 1rem; border-radius: 17px;
        box-shadow: 0 7px 20px rgba(40,52,95,.045);
    }
    div[data-testid="stMetric"] label { color: #68738c; font-weight: 650; }
    div[data-testid="stMetricValue"] { color: #302a78; font-family: 'Space Grotesk', sans-serif; }
    [data-testid="stForm"], [data-testid="stVerticalBlockBorderWrapper"] {
        border-radius: 20px;
    }
    div[data-testid="stForm"] {
        background: rgba(255,255,255,.82); border: 1px solid var(--line);
        padding: 1.2rem 1.3rem 1rem 1.3rem;
        box-shadow: 0 12px 32px rgba(40,52,95,.055);
    }
    div[data-testid="stTextInput"] input, div[data-testid="stTextArea"] textarea,
    div[data-testid="stSelectbox"] > div > div, div[data-testid="stFileUploader"] section {
        border-radius: 12px !important;
    }
    div[data-testid="stFileUploader"] section {
        background: #f6f7ff; border: 1px dashed #9b9eea;
    }
    div.stButton > button, div.stDownloadButton > button, div[data-testid="stFormSubmitButton"] button {
        border-radius: 12px; font-weight: 750; min-height: 2.75rem;
        transition: transform .15s ease, box-shadow .15s ease;
    }
    div[data-testid="stFormSubmitButton"] button, div.stDownloadButton > button {
        background: linear-gradient(105deg, #6654e8, #8b5cf6 58%, #25aebf);
        color: white; border: 0; box-shadow: 0 8px 18px rgba(102,84,232,.19);
    }
    div.stButton > button:hover, div.stDownloadButton > button:hover,
    div[data-testid="stFormSubmitButton"] button:hover {
        transform: translateY(-1px); box-shadow: 0 10px 22px rgba(102,84,232,.22);
        border-color: transparent;
    }
    div[data-testid="stDataFrame"] { border: 1px solid var(--line); border-radius: 14px; overflow: hidden; }
    [data-testid="stAlert"] { border-radius: 14px; }
    [data-testid="stStatusWidget"] { border-radius: 16px; }
    [data-testid="stSpinner"] { background: linear-gradient(110deg, rgba(102,84,232,.08), rgba(37,174,191,.10)); border: 1px solid rgba(114,87,245,.18); border-radius: 16px; padding: 1rem 1.2rem; }
    [data-testid="stSpinner"] p { color: #4338a8 !important; font-weight: 750; }
    [data-testid="stStatus"] { border: 1px solid rgba(114,87,245,.28); border-radius: 18px; overflow: hidden; box-shadow: 0 12px 35px rgba(50,49,120,.10); }
    [data-testid="stStatus"] > details > summary { background: linear-gradient(100deg, #24265d, #6650d9 68%, #24aebd); color: #fff !important; font-weight: 800; padding: 1rem 1.2rem; }
    .loader-card { text-align:center; margin:.3rem 0 .5rem; padding:1.4rem 1rem 1.2rem; border-radius:18px; background:linear-gradient(135deg,#f8f7ff,#f0fbff); border:1px solid rgba(114,87,245,.16); }
    .loader-orbit { width:76px; height:76px; margin:0 auto .65rem; display:grid; place-items:center; border-radius:50%; border:3px solid #d9d5ff; border-top-color:#7257f5; border-right-color:#35c9d6; animation:hirelens-spin 1s linear infinite; }
    .loader-core { width:55px; height:55px; display:grid; place-items:center; border-radius:50%; background:white; font-size:1.8rem; box-shadow:0 5px 18px rgba(73,65,170,.12); }
    .loader-heading { color:#25265f; font-family:'Space Grotesk',sans-serif; font-size:1.55rem; font-weight:800; letter-spacing:-.03em; }
    .loader-subtitle { color:#58627f; font-size:.92rem; margin:.35rem 0 .9rem; }
    .loader-track { max-width:340px; height:6px; margin:0 auto; border-radius:99px; background:#dedff5; overflow:hidden; }
    .loader-track-fill { height:100%; width:42%; border-radius:99px; background:linear-gradient(90deg,#7257f5,#35c9d6); animation:hirelens-progress 1.6s ease-in-out infinite; }
    .loader-note { color:#8a91a8; font-size:.75rem; margin-top:.75rem; }
    @keyframes hirelens-spin { to { transform:rotate(360deg); } }
    @keyframes hirelens-progress { 0% { transform:translateX(-120%); } 100% { transform:translateX(360%); } }


    .section-kicker { color: #7257f5; font-size: .72rem; text-transform: uppercase; letter-spacing: .13em; font-weight: 800; }
    .footer {
        margin-top: 3rem; padding: 1.35rem 1.5rem; border-radius: 20px;
        background: linear-gradient(110deg, #171f43, #343078 68%, #245b78);
        color: #f7f7ff; border: 1px solid rgba(255,255,255,.14);
    }
    .footer-title { color: #fff; font-family: 'Space Grotesk', sans-serif; font-size: 1.08rem; font-weight: 700; }
    .footer-copy { color: #d4dcff; font-size: .82rem; line-height: 1.55; margin-top: .35rem; }
    .footer-pill {
        display: inline-block; color: #e7fbff; border: 1px solid rgba(179,241,255,.28);
        background: rgba(255,255,255,.08); border-radius: 999px; padding: .35rem .65rem;
        font-size: .72rem; font-weight: 700; margin: .2rem .2rem .2rem 0;
    }
    .hint { color: #8993a8; font-size: .82rem; }
    @media (max-width: 700px) {
        .block-container { padding-left: 1rem; padding-right: 1rem; }
        .hero { border-radius: 19px; }
        div[data-testid="stForm"] { padding: .8rem; }
    }
    </style>
    """, unsafe_allow_html=True)

apply_style()

with st.sidebar:
    st.markdown("## 🎯 HireLens AI")
    st.caption("Explainable resume matching")
    page = st.radio("Navigation", ["Screening workspace", "How it works", "Responsible AI"])
    st.divider()
    demo_mode = st.toggle("Use fictional demo resumes", value=False)
    st.caption("Demo profiles contain synthetic information.")
    if st.button("Clear current results", use_container_width=True):
        for key in ["results", "resume_texts", "role_name", "run_time"]:
            st.session_state.pop(key, None)
        st.rerun()

if page == "How it works":
    st.markdown('<div class="hero"><div class="eyebrow">Product tour</div><h1>From resume pile to a reviewable shortlist.</h1><p>Compare candidate documents against role requirements, inspect evidence, and keep the recruiter in control.</p><div class="brandline">✦ HIRE SMARTER · REVIEW FAIRER · KEEP HUMANS IN CONTROL</div></div>', unsafe_allow_html=True)
    cols = st.columns(3)
    for col, num, title, desc in zip(cols, ["01", "02", "03"], ["Upload", "Match", "Review"], [
        "Add multiple PDF or DOCX resumes, or load synthetic demo profiles.",
        "Blend document similarity with must-have and preferred skill coverage.",
        "Inspect the source text and export a shortlist for human review.",
    ]):
        with col:
            st.markdown(f"### {num} · {title}")
            st.write(desc)
    st.info("This is a hackathon MVP using TF-IDF and a predefined skills dictionary. It is not a validated hiring assessment.")
elif page == "Responsible AI":
    st.markdown('<div class="hero"><div class="eyebrow">Responsible use</div><h1>AI should support hiring decisions—not make them alone.</h1><p>A match score describes document overlap with a job description, not a person’s value or future performance.</p><div class="brandline">✦ TRANSPARENT · JOB-RELATED · HUMAN-LED</div></div>', unsafe_allow_html=True)
    st.markdown("""
    - **Human review:** Do not automatically reject candidates based on a score.
    - **Job-related criteria:** Use only criteria relevant to the actual role.
    - **Fairness:** Do not score protected or irrelevant personal characteristics.
    - **Privacy:** Use synthetic resumes in public demos. Restrict access to real resumes and delete them when no longer needed.
    - **Evidence:** Verify qualifications and experience against the original resume.
    - **Testing:** Check extraction quality, missed matches, and unfair outcomes before real-world use.
    """)
    st.warning("A production system needs authentication, access controls, secure storage, retention/deletion policies, evaluation, and legal review.")
else:
    st.markdown('<div class="hero"><div class="eyebrow">✦ AI-ASSISTED TALENT DISCOVERY</div><h1>🎯 HireLens AI</h1><p>Find the fit. Explain the match. Screen resumes with flexible criteria, compare skills, and build a reviewable shortlist—with the recruiter in control.</p><div class="brandline">✦ SMART MATCHING &nbsp; · &nbsp; CLEAR EVIDENCE &nbsp; · &nbsp; HUMAN-LED HIRING</div></div>', unsafe_allow_html=True)

    with st.form("screening_form", clear_on_submit=False):
        a, b = st.columns([1.05, .95], gap="large")
        with a:
            st.markdown("#### 1 · Role details")
            role_name = st.text_input("Job title", placeholder="e.g. Junior Python Developer")
            job_description = st.text_area(
                "Job description", height=190,
                placeholder="Paste role responsibilities, qualifications, required experience and skills..."
            )
            detected = sorted(detect_skills(job_description)) if job_description else []
            must_text = st.text_input(
                "Must-have skills (comma-separated)",
                value=", ".join(detected),
                help="Required skills that should be explicitly reviewed.",
            )
            preferred_text = st.text_input(
                "Preferred skills (comma-separated)",
                placeholder="e.g. docker, aws, communication",
                help="Nice-to-have skills. These contribute less than must-have skills.",
            )
            text_weight_pct = st.slider(
                "Job-description relevance weight", 20, 80, 55, 5,
                help="Remaining score weight is split between must-have and preferred skill coverage."
            )
        with b:
            st.markdown("#### 2 · Candidate resumes")
            uploads = st.file_uploader(
                "Upload multiple PDF or DOCX files",
                type=["pdf", "docx"], accept_multiple_files=True,
                help="Text-based PDFs work best. Scanned PDFs may need OCR."
            )
            st.markdown("#### 3 · Matching preferences")
            show_low_matches = st.checkbox("Include all candidates, including lower matches", value=True)
            st.caption("The results are a review aid. Always verify the original resume.")
            st.markdown("**Included in the report**")
            st.markdown("Match score · text relevance · skill coverage · skill gaps · resume text")
        submitted = st.form_submit_button("✨ Analyze candidate matches", type="primary", use_container_width=True)

    if submitted:
        if not job_description.strip():
            st.error("Please enter a job description.")
        else:
            try:
                with st.status("🎯 HireLens AI is analyzing candidate resumes…", expanded=True) as analysis_status:
                    resumes = demo_resumes() if demo_mode else {}
                    errors = []
                    for file in uploads or []:
                        try:
                            extracted = extract_text(file)
                            if not extracted.strip():
                                errors.append(f"{file.name}: no readable text found; scanned PDFs may need OCR.")
                                continue
                            filename = file.name
                            suffix = 2
                            while filename in resumes:
                                filename = f"{suffix}_{file.name}"
                                suffix += 1
                            resumes[filename] = extracted
                        except Exception as exc:
                            errors.append(f"{file.name}: could not read ({type(exc).__name__}).")

                    if not resumes:
                        st.error("Upload at least one readable resume or enable fictional demo resumes in the sidebar.")
                    else:
                        must = parse_skill_list(must_text)
                        preferred = parse_skill_list(preferred_text)
                        results = analyze(job_description, resumes, must, preferred, text_weight_pct / 100)
                        st.session_state["results"] = results
                        st.session_state["resume_texts"] = resumes
                        st.session_state["role_name"] = role_name.strip() or "Untitled role"
                        st.session_state["run_time"] = datetime.now().strftime("%d %b %Y, %H:%M")
                    for error in errors:
                        st.warning(error)
                analysis_status.update(label="✅ HireLens AI analysis complete", state="complete", expanded=False)
            except ValueError:
                st.error("Could not compare the supplied text. Try a longer job description and readable resumes.")

    if "results" in st.session_state:
        results = st.session_state["results"].copy()
        st.divider()
        st.caption(f"SCREENING RESULTS · {st.session_state.get('run_time', '')}")
        st.header(f"Shortlist — {st.session_state.get('role_name', 'Role')}")

        x1, x2, x3, x4 = st.columns(4)
        x1.metric("Resumes analyzed", len(results))
        x2.metric("Top document match", f"{results.iloc[0]['Match score']:.1f}%")
        x3.metric("Candidates with must-have gaps", int(results["_missing"].apply(bool).sum()))
        x4.metric("Candidates to review", len(results))

        f1, f2, f3 = st.columns([1.2, 1, 1])
        with f1:
            search_term = st.text_input("Search candidates", placeholder="Filter by filename or skill")
        with f2:
            minimum = st.slider("Minimum score to display", 0, 100, 0, 5)
        with f3:
            sort_by = st.selectbox("Sort results by", ["Match score", "Text relevance", "Must-have coverage"])

        filtered = results[results["Match score"] >= minimum].copy()
        if search_term.strip():
            term = search_term.lower()
            searchable = (
                filtered["Candidate"].str.lower()
                + " " + filtered["Matching must-haves"].str.lower()
                + " " + filtered["Preferred skills found"].str.lower()
            )
            filtered = filtered[searchable.str.contains(re.escape(term), na=False)]
        if sort_by in filtered.columns:
            filtered = filtered.sort_values(sort_by, ascending=False, na_position="last")

        view_cols = [
            "Candidate", "AI recommendation", "Match score", "Text relevance",
            "Must-have coverage", "Preferred coverage",
            "Matching must-haves", "Must-haves to verify",
        ]
        st.dataframe(
            filtered[view_cols],
            use_container_width=True,
            hide_index=True,
            column_config={
                "Match score": st.column_config.ProgressColumn("Match score", min_value=0, max_value=100, format="%.1f%%"),
                "Text relevance": st.column_config.NumberColumn("Text relevance", format="%.1f%%"),
                "Must-have coverage": st.column_config.NumberColumn("Must-have coverage", format="%.1f%%"),
                "Preferred coverage": st.column_config.NumberColumn("Preferred coverage", format="%.1f%%"),
            },
        )

        if not filtered.empty:
            chosen = st.selectbox("Open candidate details", filtered["Candidate"].tolist())
            row = results[results["Candidate"] == chosen].iloc[0]
            left, right = st.columns([.9, 1.5], gap="large")
            with left:
                st.markdown("### Candidate details")
                st.metric("Document match", f"{row['Match score']:.1f}%")
                st.caption("Document relevance only — not a probability of job success.")
                st.markdown("### 🤖 AI screening recommendation")
                recommendation = row["AI recommendation"]
                if recommendation == "Recommend for interview review":
                    st.success(recommendation)
                elif recommendation == "Review gaps before proceeding":
                    st.warning(recommendation)
                else:
                    st.info(recommendation)
                st.write(row["Recommendation rationale"])
                st.caption("This is explainable triage guidance, not a final hire/no-hire decision. A human reviewer must verify the original resume and assess the candidate fairly.")
                st.write("**Matching must-have skills**")
                st.write(row["Matching must-haves"])
                st.write("**Must-have skills to verify**")
                st.write(row["Must-haves to verify"])
                st.write("**Preferred skills found**")
                st.write(row["Preferred skills found"])
            with right:
                st.markdown("### Extracted resume evidence")
                st.text_area(
                    "Check this against the original resume",
                    value=st.session_state["resume_texts"].get(chosen, ""),
                    height=290, disabled=True, label_visibility="collapsed",
                )

        export_cols = view_cols
        export_df = filtered[export_cols].copy()
        st.download_button(
            "⬇️ Download filtered shortlist (CSV)",
            data=export_df.to_csv(index=False).encode("utf-8-sig"),
            file_name="hirelens_shortlist.csv",
            mime="text/csv",
            use_container_width=True,
        )
        st.info("AI recommendations are decision-support only. Do not auto-reject candidates: resume parsing or wording may hide relevant skills. Verify evidence and make the final decision through a structured, job-related human review.")


st.markdown("""
<div class="footer" style="text-align:center">
  <div style="font-size:1.8rem;margin-bottom:.25rem">🎯</div>
  <div class="footer-title" style="font-size:1.3rem">HireLens AI <span style="color:#7de7ed">·</span> See the fit. Make the call.</div>
  <div class="footer-copy" style="max-width:760px;margin:.55rem auto 0">Thoughtful resume matching starts with clear evidence. Use every score as a guide for human review—not as an automatic hiring decision.</div>
  <div style="margin-top:.9rem">
    <span class="footer-pill">✦ Explainable matching</span>
    <span class="footer-pill">◈ Skills evidence</span>
    <span class="footer-pill">♡ Human-first hiring</span>
    <span class="footer-pill">🔒 Candidate privacy</span>
  </div>
  <div class="footer-copy" style="margin-top:1rem">Review fairly · Protect candidate data · Verify the original resume</div>
</div>
""", unsafe_allow_html=True)
